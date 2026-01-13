#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
创建测试Excel文件的Python脚本
需要安装: pip install openpyxl
"""

import sys
import io

# 设置输出编码为UTF-8
if sys.platform == 'win32':
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
except ImportError:
    print("错误: 需要安装 openpyxl 库")
    print("请运行: pip install openpyxl")
    sys.exit(1)

# 创建工作簿和工作表
wb = Workbook()
ws = wb.active
ws.title = "参赛人员"

# 设置表头
headers = ["序号", "日期", "学校", "年级", "班级", "姓名", "性别", "准考证号", "组别名称"]
ws.append(headers)

# 设置表头样式
header_fill = PatternFill(start_color="CCCCCC", end_color="CCCCCC", fill_type="solid")
header_font = Font(bold=True)

for col in range(1, len(headers) + 1):
    cell = ws.cell(row=1, column=col)
    cell.fill = header_fill
    cell.font = header_font

# 测试数据
test_data = [
    [1, "2026-1-13", "A学校", "一年级", "1班", "张三", "男", "2024001", "1组"],
    [2, "2026-1-13", "A学校", "一年级", "1班", "李四", "女", "2024002", "1组"],
    [3, "2026-1-13", "A学校", "一年级", "1班", "王五", "男", "2024003", "1组"],
    [4, "2026-1-13", "A学校", "一年级", "1班", "赵六", "男", "2024004", "2组"],
    [5, "2026-1-13", "A学校", "一年级", "1班", "钱七", "女", "2024005", "2组"],
    [6, "2026-1-13", "B学校", "一年级", "1班", "孙八", "男", "2024101", "1组"],
    [7, "2026-1-13", "B学校", "一年级", "1班", "周九", "女", "2024102", "1组"],
]

# 添加数据
for row_data in test_data:
    ws.append(row_data)

# 自动调整列宽
for col in ws.columns:
    max_length = 0
    col_letter = col[0].column_letter
    for cell in col:
        try:
            if len(str(cell.value)) > max_length:
                max_length = len(str(cell.value))
        except:
            pass
    adjusted_width = min(max_length + 2, 50)
    ws.column_dimensions[col_letter].width = adjusted_width

# 保存文件
output_file = "测试数据.xlsx"
wb.save(output_file)
print(f"测试Excel文件已创建: {output_file}")
print(f"包含 {len(test_data)} 条测试数据")

